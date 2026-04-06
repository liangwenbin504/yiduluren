#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将生成的图片嵌入到Word模板中
"""

import os
from docx import Document
from docx.shared import Inches


def embed_images_to_template(template_path, images):
    """
    将图片嵌入到Word模板中
    """
    try:
        # 打开模板文件
        doc = Document(template_path)
        
        # 在文档末尾添加图片
        doc.add_paragraph('\n')
        doc.add_paragraph('六壬相关内容')
        
        # 嵌入四柱图片
        if 'sizhu' in images:
            doc.add_paragraph('四柱信息')
            doc.add_picture(images['sizhu'], width=Inches(4))
            doc.add_paragraph('\n')
        
        # 嵌入天地盘图片
        if 'tiandi' in images:
            doc.add_paragraph('天地盘')
            doc.add_picture(images['tiandi'], width=Inches(3))
            doc.add_paragraph('\n')
        
        # 嵌入四课图片
        if 'sike' in images:
            doc.add_paragraph('四课')
            doc.add_picture(images['sike'], width=Inches(4))
            doc.add_paragraph('\n')
        
        # 嵌入三传图片
        if 'sanchuan' in images:
            doc.add_paragraph('三传')
            doc.add_picture(images['sanchuan'], width=Inches(4))
            doc.add_paragraph('\n')
        
        # 保存修改后的文档
        output_path = os.path.join(os.path.dirname(__file__), '模板_with_images.docx')
        doc.save(output_path)
        
        return output_path
    except Exception as e:
        print(f"⚠️ 嵌入图片失败: {e}")
        return None


def main():
    """
    主函数
    """
    # 模板路径
    template_path = os.path.join(os.path.dirname(__file__), '文字排版.docx')
    
    # 检查模板文件是否存在
    if not os.path.exists(template_path):
        print(f"❌ 模板文件不存在: {template_path}")
        return
    
    # 图片路径
    images = {
        'sizhu': os.path.join(os.path.dirname(__file__), 'sizhu_temp.png'),
        'tiandi': os.path.join(os.path.dirname(__file__), 'tiandi_temp.png'),
        'sike': os.path.join(os.path.dirname(__file__), 'sike_temp.png'),
        'sanchuan': os.path.join(os.path.dirname(__file__), 'sanchuan_temp.png')
    }
    
    # 检查图片是否存在
    for name, path in images.items():
        if not os.path.exists(path):
            print(f"❌ {name}图片不存在: {path}")
            return
    
    # 嵌入图片
    output_path = embed_images_to_template(template_path, images)
    
    if output_path:
        print(f"✅ 图片嵌入成功：{output_path}")
    else:
        print("❌ 图片嵌入失败")


if __name__ == '__main__':
    main()
