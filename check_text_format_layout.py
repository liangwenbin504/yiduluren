#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查文字排版.docx的布局和格式
"""

from docx import Document
from docx.oxml.ns import qn

# 打开文档
doc = Document('文字排版.docx')

print('=== 文字排版.docx 布局分析 ===')
print('段落数:', len(doc.paragraphs))
print('表格数:', len(doc.tables))
print('内联图片数:', len(doc.inline_shapes))

print('\n=== 段落内容 ===')
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print(f'段落{i}: {text[:100]}')
    else:
        print(f'段落{i}: [空]')

# 检查是否有竖排设置
print('\n=== 页面设置 ===')
section = doc.sections[0]
sectPr = section._sectPr

# 检查文字方向
textDirection = sectPr.find(qn('w:textDirection'))
if textDirection:
    print(f'文字方向: {textDirection.get(qn("w:val"))}')
else:
    print('文字方向: 默认（横排）')

# 重新复制文字排版.docx为模板-006.docx
import shutil
shutil.copy2('文字排版.docx', '模板-006.docx')
print('\n已重新复制文字排版.docx为模板-006.docx')
