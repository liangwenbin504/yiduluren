#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word文档模板导入功能
支持识别经过人工修改的模板文件
"""

import os
import re
from docx import Document

def import_template(template_path):
    """
    导入Word文档模板
    识别经过人工修改的模板文件
    """
    try:
        doc = Document(template_path)
        data = {
            'title': '',
            'four_pillars': {},
            'daliuren_data': {
                'yueJiang': '',
                'shiChen': '',
                'tianPan': {},
                'sike': [],
                'sanchuan': {}
            },
            'evaluation': ''
        }
        
        # 解析文档内容
        current_section = None
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            
            # 识别章节
            if '【标题】' in text:
                data['title'] = text.replace('【标题】', '').strip()
                current_section = 'title'
            elif '【引言】' in text:
                current_section = 'intro'
            elif '【四柱信息】' in text:
                current_section = 'four_pillars'
            elif '【天地盘】' in text:
                current_section = 'tiandi_pan'
            elif '【月将】' in text:
                # 解析月将时辰
                yuejiang_match = re.search(r'【月将】(\w+)', text)
                shichen_match = re.search(r'【时辰】(\w+)', text)
                if yuejiang_match:
                    data['daliuren_data']['yueJiang'] = yuejiang_match.group(1)
                if shichen_match:
                    data['daliuren_data']['shiChen'] = shichen_match.group(1)
            elif '【四课】' in text:
                current_section = 'si_ke'
            elif '【三传】' in text:
                current_section = 'san_chuan'
            elif '【初传】' in text:
                # 解析初传
                match = re.search(r'【初传】(\w+)(?:\((\w+)\))?', text)
                if match:
                    data['daliuren_data']['sanchuan']['初传'] = {
                        'tiangan': match.group(1)[0],
                        'dizhi': match.group(1)[1],
                        'liuqin': match.group(2) if match.group(2) else ''
                    }
            elif '【中传】' in text:
                # 解析中传
                match = re.search(r'【中传】(\w+)(?:\((\w+)\))?', text)
                if match:
                    data['daliuren_data']['sanchuan']['中传'] = {
                        'tiangan': match.group(1)[0],
                        'dizhi': match.group(1)[1],
                        'liuqin': match.group(2) if match.group(2) else ''
                    }
            elif '【末传】' in text:
                # 解析末传
                match = re.search(r'【末传】(\w+)(?:\((\w+)\))?', text)
                if match:
                    data['daliuren_data']['sanchuan']['末传'] = {
                        'tiangan': match.group(1)[0],
                        'dizhi': match.group(1)[1],
                        'liuqin': match.group(2) if match.group(2) else ''
                    }
            elif '【综合评价】' in text:
                current_section = 'evaluation'
            elif current_section == 'evaluation':
                data['evaluation'] += text + '\n'
        
        # 解析表格
        for table in doc.tables:
            # 识别四柱表格
            if len(table.rows) == 1 and len(table.columns) == 4:
                cells = table.rows[0].cells
                for i, cell in enumerate(cells):
                    text = cell.text.strip()
                    if '【年】' in text:
                        data['four_pillars']['year'] = text.replace('【年】', '').strip()
                    elif '【月】' in text:
                        data['four_pillars']['month'] = text.replace('【月】', '').strip()
                    elif '【日】' in text:
                        data['four_pillars']['day'] = text.replace('【日】', '').strip()
                    elif '【时】' in text:
                        data['four_pillars']['hour'] = text.replace('【时】', '').strip()
            
            # 识别天地盘表格
            elif len(table.rows) == 4 and len(table.columns) == 4:
                # 定义地盘顺序
                di_pan_order = [
                    ['巳', '午', '未', '申'],
                    ['辰', '', '', '酉'],
                    ['卯', '', '', '戌'],
                    ['寅', '丑', '子', '亥']
                ]
                
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        text = cell.text.strip()
                        if text:
                            # 解析地盘和天盘
                            di_match = re.search(r'【地盘】(\w+)', text)
                            tian_match = re.search(r'【天盘】(\w+)', text)
                            if di_match and tian_match:
                                di_zhi = di_match.group(1)
                                tian_zhi = tian_match.group(1)
                                data['daliuren_data']['tianPan'][di_zhi] = tian_zhi
            
            # 识别四课表格
            elif len(table.rows) == 2 and len(table.columns) == 4:
                sike = []
                top_row = table.rows[0]
                bottom_row = table.rows[1]
                
                for i in range(4):
                    top_text = top_row.cells[i].text.strip()
                    bottom_text = bottom_row.cells[i].text.strip()
                    
                    top = ''
                    bottom = ''
                    
                    if '【上】' in top_text:
                        top = top_text.replace('【上】', '').strip()
                    if '【下】' in bottom_text:
                        bottom = bottom_text.replace('【下】', '').strip()
                    
                    if top or bottom:
                        sike.append({'top': top, 'bottom': bottom})
                
                data['daliuren_data']['sike'] = sike
        
        # 清理评价内容
        data['evaluation'] = data['evaluation'].strip()
        
        return {
            'success': True,
            'data': data
        }
        
    except Exception as e:
        import traceback
        return {
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }

def test_import():
    """
    测试模板导入功能
    """
    template_path = os.path.join(os.path.dirname(__file__), 'template.docx')
    result = import_template(template_path)
    
    if result['success']:
        print("✅ 模板导入成功")
        print("标题:", result['data']['title'])
        print("四柱:", result['data']['four_pillars'])
        print("月将:", result['data']['daliuren_data']['yueJiang'])
        print("时辰:", result['data']['daliuren_data']['shiChen'])
        print("天地盘:", result['data']['daliuren_data']['tianPan'])
        print("四课:", result['data']['daliuren_data']['sike'])
        print("三传:", result['data']['daliuren_data']['sanchuan'])
        print("评价:", result['data']['evaluation'][:100] + "..." if len(result['data']['evaluation']) > 100 else result['data']['evaluation'])
    else:
        print("❌ 模板导入失败:", result['error'])

if __name__ == '__main__':
    import sys
    if len(sys.argv) > 1:
        template_path = sys.argv[1]
        result = import_template(template_path)
        if result['success']:
            print("✅ 模板导入成功")
            print("标题:", result['data']['title'])
            print("四柱:", result['data']['four_pillars'])
            print("月将:", result['data']['daliuren_data']['yueJiang'])
            print("时辰:", result['data']['daliuren_data']['shiChen'])
            print("天地盘:", result['data']['daliuren_data']['tianPan'])
            print("四课:", result['data']['daliuren_data']['sike'])
            print("三传:", result['data']['daliuren_data']['sanchuan'])
            print("评价:", result['data']['evaluation'][:100] + "..." if len(result['data']['evaluation']) > 100 else result['data']['evaluation'])
        else:
            print("❌ 模板导入失败:", result['error'])
    else:
        test_import()
