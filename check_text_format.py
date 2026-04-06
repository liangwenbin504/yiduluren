#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查文字排版.docx的内容和结构
"""

from docx import Document

# 打开文档
doc = Document('文字排版.docx')

print('=== 文字排版.docx 结构分析 ===')
print('段落数:', len(doc.paragraphs))
print('表格数:', len(doc.tables))
print('内联图片数:', len(doc.inline_shapes))

print('\n=== 段落内容 ===')
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print(f'段落{i}: {text[:150]}...' if len(text) > 150 else f'段落{i}: {text}')
    else:
        print(f'段落{i}: [空]')

# 复制为新的模板文件
import shutil
shutil.copy2('文字排版.docx', '模板-006.docx')
print('\n已将文字排版.docx复制为模板-006.docx')
