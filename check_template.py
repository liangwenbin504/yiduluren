#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查模板-006的内容和结构
"""

from docx import Document

# 打开模板文件
doc = Document('模板-006.docx')

print('=== 模板-006.docx 结构分析 ===')
print('段落数:', len(doc.paragraphs))
print('表格数:', len(doc.tables))
print('内联图片数:', len(doc.inline_shapes))

print('\n=== 段落内容 ===')
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print(f'段落{i}: {text[:100]}...' if len(text) > 100 else f'段落{i}: {text}')
    else:
        print(f'段落{i}: [空]')

print('\n=== 表格内容 ===')
for i, table in enumerate(doc.tables):
    print(f'表格{i}: {len(table.rows)}行 × {len(table.columns)}列')
    for row_idx, row in enumerate(table.rows):
        cells_text = []
        for cell in row.cells:
            cell_text = cell.text.strip()
            cells_text.append(cell_text[:20] + '...' if len(cell_text) > 20 else cell_text)
        print(f'  行{row_idx}: {cells_text}')

print('\n=== 内联图片 ===')
for i, shape in enumerate(doc.inline_shapes):
    print(f'图片{i}: 类型={shape.type}, 宽度={shape.width}, 高度={shape.height}')
