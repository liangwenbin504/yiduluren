#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查导出的文档格式、布局和内容逻辑
"""

import os
import glob
from docx import Document
from docx.shared import Pt

def check_document(file_path):
    """检查文档格式、布局和内容逻辑"""
    print(f"\n{'='*80}")
    print(f"检查文档：{os.path.basename(file_path)}")
    print(f"{'='*80}\n")
    
    try:
        doc = Document(file_path)
        
        # 1. 检查页面设置
        print("【1. 页面设置检查】")
        for i, section in enumerate(doc.sections):
            print(f"  节 {i+1}:")
            print(f"    - 页面宽度：{section.page_width.cm:.2f} cm")
            print(f"    - 页面高度：{section.page_height.cm:.2f} cm")
            print(f"    - 上边距：{section.top_margin.cm:.2f} cm")
            print(f"    - 下边距：{section.bottom_margin.cm:.2f} cm")
            print(f"    - 左边距：{section.left_margin.cm:.2f} cm")
            print(f"    - 右边距：{section.right_margin.cm:.2f} cm")
            
            # 检查页面背景
            sectPr = section._sectPr
            bg = sectPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bg')
            if bg is not None:
                bg_color = bg.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color')
                print(f"    - 页面背景颜色：#{bg_color}")
            else:
                print("    - ⚠️ 页面背景颜色：未设置")
            
            # 检查页面边框
            pgBorders = sectPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgBorders')
            if pgBorders is not None:
                print("    - 页面边框：已设置")
                # 检查每个边框
                for border_name in ['top', 'left', 'bottom', 'right']:
                    border = pgBorders.find(f'.//{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{border_name}')
                    if border is not None:
                        border_val = border.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                        border_sz = border.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
                        print(f"      - {border_name}边框：{border_val}, 宽度：{border_sz}")
            else:
                print("    - ⚠️ 页面边框：未设置")
        
        # 2. 检查段落内容
        print("\n【2. 段落内容检查】")
        for i, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                print(f"  段落 {i+1}:")
                print(f"    - 内容：{paragraph.text[:50]}..." if len(paragraph.text) > 50 else f"    - 内容：{paragraph.text}")
                
                # 检查文本方向
                pPr = paragraph._element.get_or_add_pPr()
                textDirection = pPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}textDirection')
                if textDirection is not None:
                    direction = textDirection.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                    print(f"    - 文本方向：{direction} (tbRl表示从右到左竖排)")
                else:
                    print("    - ⚠️ 文本方向：未设置（默认为横排）")
                
                # 检查字体
                for run in paragraph.runs:
                    if run.text.strip():
                        font_name = run.font.name
                        font_size = run.font.size.pt if run.font.size else "未设置"
                        
                        # 检查东亚字体
                        rPr = run._element.get_or_add_rPr()
                        rFonts = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                        east_asia_font = None
                        if rFonts is not None:
                            east_asia_font = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
                        
                        print(f"    - 字体：{font_name}, 东亚字体：{east_asia_font}, 大小：{font_size}pt")
                        break
        
        # 3. 检查图片
        print("\n【3. 图片检查】")
        image_count = 0
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run._element.xpath('.//a:blip'):
                    image_count += 1
                    print(f"  - 发现图片 {image_count}")
        
        if image_count == 0:
            print("  - ⚠️ 未发现图片")
        
        # 4. 检查表格
        print("\n【4. 表格检查】")
        table_count = len(doc.tables)
        print(f"  - 表格数量：{table_count}")
        
        # 5. 检查内容逻辑
        print("\n【5. 内容逻辑检查】")
        content = '\n'.join([p.text for p in doc.paragraphs])
        
        # 检查必要的内容
        required_content = [
            ('标题', '课单'),
            ('开头固定文本', '伏以，天地定位'),
            ('仙命信息', '仙命'),
            ('山向信息', '山向'),
            ('四柱信息', '择取岁次'),
            ('六壬课格', '六壬课格'),
            ('批语内容', '批语'),
            ('避忌事宜', '避忌事宜'),
            ('地理师落款', '沐手谨择'),
            ('日期落款', '榖旦 勒石')
        ]
        
        for name, keyword in required_content:
            if keyword in content:
                print(f"  ✅ {name}：已包含")
            else:
                print(f"  ⚠️ {name}：未包含")
        
        # 6. 检查格式问题
        print("\n【6. 格式问题检查】")
        issues = []
        
        # 检查页面背景
        sectPr = doc.sections[0]._sectPr
        bg = sectPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bg')
        if bg is None:
            issues.append("页面背景未设置为淡黄色")
        
        # 检查页面边框
        pgBorders = sectPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgBorders')
        if pgBorders is None:
            issues.append("页面边框未设置为双线设计")
        
        # 检查文本方向
        vertical_count = 0
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                pPr = paragraph._element.get_or_add_pPr()
                textDirection = pPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}textDirection')
                if textDirection is not None:
                    vertical_count += 1
        
        if vertical_count == 0:
            issues.append("所有段落均为横排，未设置为竖排")
        
        # 检查字体
        font_issues = []
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    rPr = run._element.get_or_add_rPr()
                    rFonts = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                    if rFonts is not None:
                        east_asia_font = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
                        if east_asia_font != '李白赋诗游龙惊鸿':
                            font_issues.append(f"字体不是'李白赋诗游龙惊鸿'：{east_asia_font}")
        
        if font_issues:
            issues.append(f"部分文本字体不正确：{', '.join(set(font_issues[:3]))}")
        
        if issues:
            print("  发现以下问题：")
            for i, issue in enumerate(issues, 1):
                print(f"    {i}. {issue}")
        else:
            print("  ✅ 未发现格式问题")
        
        print(f"\n{'='*80}")
        print("检查完成")
        print(f"{'='*80}\n")
        
        return issues
        
    except Exception as e:
        import traceback
        print(f"❌ 检查文档失败：{e}")
        print(traceback.format_exc())
        return []

if __name__ == '__main__':
    # 找到最新的导出文档
    output_dir = os.path.dirname(__file__)
    docx_files = glob.glob(os.path.join(output_dir, '择日课单_*.docx'))
    
    if docx_files:
        # 按修改时间排序，获取最新的文件
        latest_file = max(docx_files, key=os.path.getmtime)
        check_document(latest_file)
    else:
        print("❌ 未找到导出的文档")
