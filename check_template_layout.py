#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
检查模板-006.docx的布局和内容
"""

from docx import Document

# 打开模板文件
doc = Document('模板-006.docx')

print('=== 模板-006.docx 布局分析 ===')
print('段落数:', len(doc.paragraphs))
print('表格数:', len(doc.tables))
print('内联图片数:', len(doc.inline_shapes))

print('\n=== 前10个段落内容 ===')
for i, para in enumerate(doc.paragraphs[:10]):
    text = para.text.strip()
    if text:
        print(f'段落{i}: {text[:150]}')
    else:
        print(f'段落{i}: [空]')

print('\n=== 表格信息 ===')
for i, table in enumerate(doc.tables):
    print(f'表格{i}: {len(table.rows)}行 × {len(table.columns)}列')

print('\n=== 内联图片 ===')
for i, shape in enumerate(doc.inline_shapes):
    print(f'图片{i}: 类型={shape.type}, 宽度={shape.width}, 高度={shape.height}')
