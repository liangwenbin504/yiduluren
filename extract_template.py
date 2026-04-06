# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm

doc = Document('文字排版.docx')

output = []

output.append('='*70)
output.append('【文字排版.docx 模板完整内容】')
output.append('='*70)

section = doc.sections[0]
output.append('\n【页面设置】')
output.append(f'页面宽度: {section.page_width.cm:.2f} cm')
output.append(f'页面高度: {section.page_height.cm:.2f} cm')
output.append(f'上边距: {section.top_margin.cm:.2f} cm')
output.append(f'下边距: {section.bottom_margin.cm:.2f} cm')
output.append(f'左边距: {section.left_margin.cm:.2f} cm')
output.append(f'右边距: {section.right_margin.cm:.2f} cm')
if section.page_width.cm > section.page_height.cm:
    output.append('纸张方向: 横向')
else:
    output.append('纸张方向: 纵向')

output.append(f'\n【完整文档内容】共 {len(doc.paragraphs)} 个段落')
output.append('-'*70)

for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        output.append(f'\n[段落 {i+1}]')
        output.append(f'对齐方式: {para.alignment}')
        output.append(f'内容:')
        output.append(para.text)
        output.append(f'格式详情:')
        for j, run in enumerate(para.runs):
            font_size = f'{run.font.size.pt}pt' if run.font.size else '默认'
            font_name = run.font.name or '默认'
            bold = '粗体' if run.font.bold else '常规'
            text_preview = run.text[:30] if len(run.text) > 30 else run.text
            output.append(f'  Run {j+1}: {text_preview}... | 字体:{font_name} | 大小:{font_size} | {bold}')

if doc.tables:
    output.append(f'\n【表格分析】共 {len(doc.tables)} 个表格')
    for i, table in enumerate(doc.tables):
        output.append(f'\n表格 {i+1}: {len(table.rows)} 行 x {len(table.columns)} 列')
        for row in table.rows:
            row_text = [cell.text[:20] for cell in row.cells]
            output.append(f'  {row_text}')

result = '\n'.join(output)
print(result)

with open('模板内容分析.txt', 'w', encoding='utf-8') as f:
    f.write(result)

print('\n\n内容已保存到: 模板内容分析.txt')
