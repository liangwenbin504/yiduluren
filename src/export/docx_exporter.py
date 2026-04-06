"""
DOCX 导出功能模块
实现六壬盘式的美观导出
"""

import io
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class LiuRenExporter:
    """六壬盘式导出器"""
    
    def __init__(self):
        # 颜色配置（与主界面一致）
        self.colors = {
            'bg_tianjiang': (255, 228, 196),  # 天将背景 - 浅橙色
            'bg_dizhi': (176, 208, 240),      # 地支背景 - 中等蓝色
            'text_dizhi': (0, 0, 255),        # 地支文字 - 纯蓝色
            'text_tian': (205, 92, 92),       # 天盘文字 - 深红色
            'text_tianjiang': (139, 69, 19),  # 天将文字 - 棕色
            'border': (139, 115, 85),         # 边框 - 古铜色
            'header': (139, 0, 0),            # 标题 - 深红色
            'bg_sike_top': (255, 182, 193),   # 四课上神背景 - 浅红色
            'bg_sike_bottom': (135, 206, 235), # 四课下神背景 - 天蓝色
            'bg_sanchuan': (152, 251, 152),   # 三传背景 - 浅绿色
            'text_common': (44, 62, 80),      # 通用文字颜色
        }
        
        # 布局参数
        self.cell_size = 55
        self.margin_left = 100
        self.margin_top = 60
        
    def generate_tiandi_pan_image(self, tiandi_pan, tianjiang_list=None):
        """
        生成天地盘图片
        :param tiandi_pan: 天地盘数据
        :param tianjiang_list: 天将列表
        :return: PIL Image 对象
        """
        # 创建图片（宽度 700，高度 550）
        img = Image.new('RGB', (700, 550), (250, 240, 230))  # 米色背景
        draw = ImageDraw.Draw(img)
        
        # 尝试加载字体
        try:
            font = ImageFont.truetype('msyh.ttf', 12)
            font_bold = ImageFont.truetype('msyhbd.ttf', 14)
            font_title = ImageFont.truetype('msyhbd.ttf', 20)
        except:
            # 如果字体不存在，使用默认字体
            font = ImageFont.load_default()
            font_bold = ImageFont.load_default()
            font_title = ImageFont.load_default()
        
        # 地盘位置（方形布局）
        dizhi_positions = {
            # 上排（4 个，从左到右）
            '巳': (0, 0), '午': (1, 0), '未': (2, 0), '申': (3, 0),
            # 左列（2 个，从上到下）
            '辰': (0, 1), '卯': (0, 2),
            # 右列（2 个，从上到下，左移一格）
            '酉': (3, 1), '戌': (3, 2),
            # 下排（4 个，从右到左）
            '亥': (3, 3), '子': (2, 3), '丑': (1, 3), '寅': (0, 3),
        }
        
        # 天将位置
        outer_positions = {
            '子': (0, -1), '丑': (1, -1), '寅': (2, -1),
            '卯': (4, 0), '辰': (4, 1), '巳': (4, 2),
            '午': (3, 4), '未': (2, 4), '申': (1, 4),
            '酉': (-1, 3), '戌': (-1, 2), '亥': (-1, 1),
        }
        
        # 1. 绘制天将（外圈）
        if tianjiang_list:
            # 创建天将映射：天盘地支 -> 天将名
            tianjiang_map = {}
            for tj in tianjiang_list:
                tian_zhi = tj['地支']
                tianjiang_name = tj['天将']
                short_name = self._get_tianjiang_short(tianjiang_name)
                tianjiang_map[tian_zhi] = short_name
            
            # 创建地盘->天盘的逆映射：天盘地支 -> 地盘地支
            tian_to_di = {}
            for di_zhi, tian_zhi in tiandi_pan['天地对应'].items():
                tian_to_di[tian_zhi] = di_zhi
            
            # 绘制每个天将
            for tian_zhi, short_name in tianjiang_map.items():
                di_zhi = tian_to_di.get(tian_zhi)
                if not di_zhi:
                    continue
                
                di_pos = dizhi_positions.get(di_zhi)
                if not di_pos:
                    continue
                
                col, row = di_pos
                tj_pos = outer_positions.get(di_zhi)
                if not tj_pos:
                    continue
                
                tj_col, tj_row = tj_pos
                x = self.margin_left + tj_col * self.cell_size
                y = self.margin_top + tj_row * self.cell_size
                
                # 绘制背景格
                draw.rectangle(
                    [x, y, x + self.cell_size, y + self.cell_size],
                    fill=self.colors['bg_tianjiang'],
                    outline=self.colors['border'],
                    width=1
                )
                
                # 绘制天将文字
                draw.text(
                    (x + self.cell_size//2, y + self.cell_size//2),
                    short_name,
                    font=font,
                    fill=self.colors['text_tianjiang'],
                    anchor='mm'
                )
        
        # 2. 绘制天地盘（内圈）
        for dizhi, (col, row) in dizhi_positions.items():
            x = self.margin_left + col * self.cell_size
            y = self.margin_top + row * self.cell_size
            
            # 绘制背景格
            draw.rectangle(
                [x, y, x + self.cell_size, y + self.cell_size],
                fill=self.colors['bg_dizhi'],
                outline=self.colors['border'],
                width=1
            )
            
            # 获取天盘地支
            tian_zhi = tiandi_pan['天地对应'].get(dizhi, '')
            
            # 绘制天盘（红色，上方）
            if tian_zhi:
                draw.text(
                    (x + self.cell_size//2, y + self.cell_size//2 - 12),
                    tian_zhi,
                    font=font_bold,
                    fill=self.colors['text_tian'],
                    anchor='mm'
                )
            
            # 绘制地盘（蓝色，下方）
            draw.text(
                (x + self.cell_size//2, y + self.cell_size//2 + 12),
                dizhi,
                font=font_bold,
                fill=self.colors['text_dizhi'],
                anchor='mm'
            )
        
        # 3. 绘制标题
        draw.text(
            (self.margin_left + 2 * self.cell_size, self.margin_top - 40),
            '天地盘',
            font=font_title,
            fill=self.colors['header'],
            anchor='mm'
        )
        
        # 4. 绘制图例说明
        draw.text(
            (self.margin_left + 2 * self.cell_size, self.margin_top - 15),
            '▲天盘 (红)  ▼地盘 (蓝)',
            font=font,
            fill=(102, 102, 102),
            anchor='mm'
        )
        
        return img
    
    def generate_si_ke_image(self, si_ke, ri_zhu):
        """
        生成四课图片
        :param si_ke: 四课数据
        :param ri_zhu: 日柱
        :return: PIL Image 对象
        """
        # 创建图片（宽度 400，高度 300）
        img = Image.new('RGB', (400, 300), (250, 240, 230))  # 米色背景
        draw = ImageDraw.Draw(img)
        
        # 尝试加载字体
        try:
            font = ImageFont.truetype('msyh.ttf', 10)
            font_bold = ImageFont.truetype('msyhbd.ttf', 14)
            font_title = ImageFont.truetype('msyhbd.ttf', 16)
        except:
            font = ImageFont.load_default()
            font_bold = ImageFont.load_default()
            font_title = ImageFont.load_default()
        
        cell_size = 60
        x_start = 50
        y_start = 50
        
        # 标题
        draw.text(
            (x_start + 150, y_start - 20),
            '四课',
            font=font_title,
            fill=self.colors['header'],
            anchor='mm'
        )
        
        # 四课布局：从右到左排列
        ke_positions = [3, 2, 1, 0]
        offset_x = 80
        
        for i, ke in enumerate(si_ke):
            col = ke_positions[i]
            x = x_start + col * cell_size + offset_x
            
            # 课名标签
            ke_name = f"第{['一','二','三','四'][i]}课"
            draw.rectangle(
                [x, y_start - 25, x + cell_size, y_start - 5],
                fill=(255, 215, 0),  # 金色
                outline=self.colors['border'],
                width=1
            )
            draw.text(
                (x + cell_size//2, y_start - 15),
                ke_name,
                font=font,
                fill=self.colors['text_common'],
                anchor='mm'
            )
            
            # 上神
            draw.rectangle(
                [x, y_start, x + cell_size, y_start + cell_size],
                fill=self.colors['bg_sike_top'],
                outline=self.colors['border'],
                width=2
            )
            draw.text(
                (x + cell_size//2, y_start + cell_size//2),
                ke['top'],
                font=font_bold,
                fill=self.colors['text_common'],
                anchor='mm'
            )
            
            # 下神
            draw.rectangle(
                [x, y_start + cell_size, x + cell_size, y_start + 2*cell_size],
                fill=self.colors['bg_sike_bottom'],
                outline=self.colors['border'],
                width=2
            )
            draw.text(
                (x + cell_size//2, y_start + cell_size*1.5),
                ke['bottom'],
                font=font_bold,
                fill=self.colors['text_common'],
                anchor='mm'
            )
        
        # 标注"四课"竖排标题
        draw.text(
            (x_start + 10, y_start + cell_size),
            '四课',
            font=font_bold,
            fill=self.colors['header'],
            anchor='mm'
        )
        
        return img
    
    def generate_san_chuan_image(self, san_chuan):
        """
        生成三传图片
        :param san_chuan: 三传数据
        :return: PIL Image 对象
        """
        # 创建图片（宽度 400，高度 250）
        img = Image.new('RGB', (400, 250), (250, 240, 230))  # 米色背景
        draw = ImageDraw.Draw(img)
        
        # 尝试加载字体
        try:
            font = ImageFont.truetype('msyh.ttf', 10)
            font_bold = ImageFont.truetype('msyhbd.ttf', 14)
            font_title = ImageFont.truetype('msyhbd.ttf', 16)
        except:
            font = ImageFont.load_default()
            font_bold = ImageFont.load_default()
            font_title = ImageFont.load_default()
        
        cell_size = 50
        x_start = 50
        y_start = 50
        
        # 标题
        draw.text(
            (x_start + 150, y_start - 20),
            '三传',
            font=font_title,
            fill=self.colors['header'],
            anchor='mm'
        )
        
        chuan_list = san_chuan['三传']
        chuan_names = ['初传', '中传', '末传']
        
        # 三传竖向排列
        for i, chuan in enumerate(chuan_list):
            if chuan:
                y = y_start + i * cell_size
                x = x_start + 80
                
                # 传名标签
                draw.text(
                    (x - 40, y + cell_size//2),
                    chuan_names[i],
                    font=font_bold,
                    fill=self.colors['header'],
                    anchor='mm'
                )
                
                # 传
                draw.rectangle(
                    [x, y, x + cell_size*1.5, y + cell_size],
                    fill=self.colors['bg_sanchuan'],
                    outline=self.colors['border'],
                    width=2
                )
                draw.text(
                    (x + cell_size*0.75, y + cell_size//2),
                    chuan,
                    font=font_bold,
                    fill=self.colors['text_common'],
                    anchor='mm'
                )
        
        # 课体和起法
        draw.text(
            (x_start + 280, y_start + cell_size//2),
            f"课体：{san_chuan['课体']}",
            font=font_bold,
            fill=self.colors['header'],
            anchor='mm'
        )
        draw.text(
            (x_start + 280, y_start + cell_size*1.5),
            f"起法：{san_chuan['起法']}",
            font=font_bold,
            fill=self.colors['header'],
            anchor='mm'
        )
        
        return img
    
    def export_to_docx(self, result, file_path):
        """
        导出到 DOCX 文档
        :param result: 六壬排盘结果
        :param file_path: 导出文件路径
        """
        # 创建文档
        doc = Document()
        
        # 设置标题
        doc.add_heading('大六壬排盘结果', level=1)
        
        # 基本信息
        info_paragraph = doc.add_paragraph()
        info_paragraph.add_run('基本信息').bold = True
        info_paragraph.add_run('\n')
        info_paragraph.add_run(f'月将：{result["月将"]}\n')
        info_paragraph.add_run(f'占时：{result["占时"]}\n')
        info_paragraph.add_run(f'日柱：{result["日柱"]}\n')
        
        # 天地盘图片
        tiandi_pan_img = self.generate_tiandi_pan_image(result['天地盘'], result.get('天将'))
        img_bytes = io.BytesIO()
        tiandi_pan_img.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        
        tiandi_paragraph = doc.add_paragraph()
        tiandi_paragraph.add_run('天地盘').bold = True
        tiandi_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(img_bytes, width=Inches(6.0))
        
        # 四课图片
        si_ke_img = self.generate_si_ke_image(result['四课'], result['日柱'])
        img_bytes = io.BytesIO()
        si_ke_img.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        
        si_ke_paragraph = doc.add_paragraph()
        si_ke_paragraph.add_run('四课').bold = True
        si_ke_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(img_bytes, width=Inches(4.0))
        
        # 三传图片
        san_chuan_img = self.generate_san_chuan_image(result['三传'])
        img_bytes = io.BytesIO()
        san_chuan_img.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        
        san_chuan_paragraph = doc.add_paragraph()
        san_chuan_paragraph.add_run('三传').bold = True
        san_chuan_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(img_bytes, width=Inches(4.0))
        
        # 禄马贵信息
        lu_ma_gui_paragraph = doc.add_paragraph()
        lu_ma_gui_paragraph.add_run('禄马贵').bold = True
        lu_ma_gui_paragraph.add_run('\n')
        lu_ma_gui_paragraph.add_run(f'禄：{result["禄马贵"]["禄"]}\n')
        lu_ma_gui_paragraph.add_run(f'驿马：{result["禄马贵"]["驿马"]}\n')
        lu_ma_gui_paragraph.add_run(f'贵人：{"、".join(result["禄马贵"]["贵人"])}\n')
        
        # 保存文档
        doc.save(file_path)
    
    def _get_tianjiang_short(self, name):
        """获取天将简称"""
        mapping = {
            '贵人': '贵', '螣蛇': '蛇', '朱雀': '雀', '六合': '合',
            '勾陈': '勾', '青龙': '龙', '天空': '空', '白虎': '虎',
            '太常': '常', '玄武': '玄', '太阴': '阴', '天后': '后'
        }
        return mapping.get(name, name[0])
