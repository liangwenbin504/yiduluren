"""
文档格式重新设计模块
实现文档格式的自定义设置和模板管理
"""

from docx import Document
from docx.shared import RGBColor, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION
import json
import os


class DocxFormatter:
    """DOCX 文档格式编辑器"""
    
    def __init__(self):
        """初始化文档格式编辑器"""
        # 默认格式设置
        self.default_format = {
            'font': {
                'name': '微软雅黑',
                'size': 12,
                'color': (0, 0, 0),  # 黑色
                'bold': False,
                'italic': False
            },
            'heading': {
                'h1': {
                    'font': '微软雅黑',
                    'size': 20,
                    'color': (139, 0, 0),  # 深红色
                    'bold': True
                },
                'h2': {
                    'font': '微软雅黑',
                    'size': 16,
                    'color': (139, 0, 0),
                    'bold': True
                }
            },
            'paragraph': {
                'alignment': 'left',  # left, center, right, justify
                'space_before': 0,
                'space_after': 10,
                'line_spacing': 1.5
            },
            'image': {
                'width': 6.0,  # 英寸
                'alignment': 'center',
                'wrap_style': 'inline'  # inline, square, tight, behind, in_front
            },
            'header': {
                'text': '仪度六壬择日系统',
                'font': '微软雅黑',
                'size': 10,
                'color': (0, 0, 0),
                'alignment': 'center'
            },
            'footer': {
                'text': '第 {page} 页，共 {pages} 页',
                'font': '微软雅黑',
                'size': 9,
                'color': (102, 102, 102),
                'alignment': 'center'
            },
            'page': {
                'orientation': 'portrait',  # portrait, landscape
                'margin_top': 1.0,
                'margin_bottom': 1.0,
                'margin_left': 1.0,
                'margin_right': 1.0
            }
        }
        
        # 当前格式设置
        self.current_format = self.default_format.copy()
        
        # 模板目录
        self.template_dir = os.path.join(os.path.dirname(__file__), 'templates')
        if not os.path.exists(self.template_dir):
            os.makedirs(self.template_dir)
    
    def load_template(self, template_name):
        """
        加载模板
        :param template_name: 模板名称
        :return: 是否加载成功
        """
        template_path = os.path.join(self.template_dir, f'{template_name}.json')
        if os.path.exists(template_path):
            try:
                with open(template_path, 'r', encoding='utf-8') as f:
                    self.current_format = json.load(f)
                return True
            except Exception as e:
                print(f"加载模板失败: {e}")
                return False
        return False
    
    def save_template(self, template_name):
        """
        保存模板
        :param template_name: 模板名称
        :return: 是否保存成功
        """
        template_path = os.path.join(self.template_dir, f'{template_name}.json')
        try:
            with open(template_path, 'w', encoding='utf-8') as f:
                json.dump(self.current_format, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            print(f"保存模板失败: {e}")
            return False
    
    def list_templates(self):
        """
        列出所有可用模板
        :return: 模板名称列表
        """
        templates = []
        if os.path.exists(self.template_dir):
            for file in os.listdir(self.template_dir):
                if file.endswith('.json'):
                    templates.append(file[:-5])  # 移除 .json 扩展名
        return templates
    
    def reset_format(self):
        """
        重置为默认格式
        """
        self.current_format = self.default_format.copy()
    
    def update_format(self, format_dict):
        """
        更新格式设置
        :param format_dict: 格式设置字典
        """
        def update_recursive(current, update):
            for key, value in update.items():
                if isinstance(value, dict) and key in current:
                    update_recursive(current[key], value)
                else:
                    current[key] = value
        
        update_recursive(self.current_format, format_dict)
    
    def apply_format(self, doc):
        """
        应用格式到文档
        :param doc: Document 对象
        """
        # 设置页面格式
        for section in doc.sections:
            # 页面方向
            if self.current_format['page']['orientation'] == 'landscape':
                section.orientation = WD_ORIENTATION.LANDSCAPE
            else:
                section.orientation = WD_ORIENTATION.PORTRAIT
            
            # 页边距
            section.top_margin = Inches(self.current_format['page']['margin_top'])
            section.bottom_margin = Inches(self.current_format['page']['margin_bottom'])
            section.left_margin = Inches(self.current_format['page']['margin_left'])
            section.right_margin = Inches(self.current_format['page']['margin_right'])
            
            # 设置页眉
            header = section.header
            header_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            header_paragraph.text = self.current_format['header']['text']
            run = header_paragraph.runs[0]
            run.font.name = self.current_format['header']['font']
            run.font.size = Pt(self.current_format['header']['size'])
            run.font.color.rgb = RGBColor(*self.current_format['header']['color'])
            
            # 设置页眉对齐方式
            align_map = {
                'left': WD_ALIGN_PARAGRAPH.LEFT,
                'center': WD_ALIGN_PARAGRAPH.CENTER,
                'right': WD_ALIGN_PARAGRAPH.RIGHT
            }
            header_paragraph.alignment = align_map.get(self.current_format['header']['alignment'], WD_ALIGN_PARAGRAPH.LEFT)
            
            # 设置页脚
            footer = section.footer
            footer_paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            footer_paragraph.text = self.current_format['footer']['text']
            run = footer_paragraph.runs[0]
            run.font.name = self.current_format['footer']['font']
            run.font.size = Pt(self.current_format['footer']['size'])
            run.font.color.rgb = RGBColor(*self.current_format['footer']['color'])
            footer_paragraph.alignment = align_map.get(self.current_format['footer']['alignment'], WD_ALIGN_PARAGRAPH.CENTER)
        
        # 设置正文格式
        for paragraph in doc.paragraphs:
            # 设置段落对齐方式
            paragraph.alignment = align_map.get(self.current_format['paragraph']['alignment'], WD_ALIGN_PARAGRAPH.LEFT)
            
            # 设置段落间距
            paragraph.space_before = Pt(self.current_format['paragraph']['space_before'])
            paragraph.space_after = Pt(self.current_format['paragraph']['space_after'])
            
            # 设置行间距
            paragraph.line_spacing = self.current_format['paragraph']['line_spacing']
            
            # 设置字体格式
            for run in paragraph.runs:
                run.font.name = self.current_format['font']['name']
                run.font.size = Pt(self.current_format['font']['size'])
                run.font.color.rgb = RGBColor(*self.current_format['font']['color'])
                run.font.bold = self.current_format['font']['bold']
                run.font.italic = self.current_format['font']['italic']
        
        # 设置标题格式
        for paragraph in doc.paragraphs:
            if paragraph.style.name.startswith('Heading'):
                level = int(paragraph.style.name[-1])
                if level == 1 and 'h1' in self.current_format['heading']:
                    for run in paragraph.runs:
                        run.font.name = self.current_format['heading']['h1']['font']
                        run.font.size = Pt(self.current_format['heading']['h1']['size'])
                        run.font.color.rgb = RGBColor(*self.current_format['heading']['h1']['color'])
                        run.font.bold = self.current_format['heading']['h1']['bold']
                elif level == 2 and 'h2' in self.current_format['heading']:
                    for run in paragraph.runs:
                        run.font.name = self.current_format['heading']['h2']['font']
                        run.font.size = Pt(self.current_format['heading']['h2']['size'])
                        run.font.color.rgb = RGBColor(*self.current_format['heading']['h2']['color'])
                        run.font.bold = self.current_format['heading']['h2']['bold']
        
        # 设置图片格式（暂时跳过，确保基本功能正常工作）
        # 图片处理需要根据具体的 python-docx 版本进行调整
        pass
    
    def import_document(self, file_path):
        """
        导入文档
        :param file_path: 文档路径
        :return: Document 对象
        """
        try:
            doc = Document(file_path)
            return doc
        except Exception as e:
            print(f"导入文档失败: {e}")
            return None
    
    def export_document(self, doc, file_path):
        """
        导出文档
        :param doc: Document 对象
        :param file_path: 导出路径
        :return: 是否导出成功
        """
        try:
            doc.save(file_path)
            return True
        except Exception as e:
            print(f"导出文档失败: {e}")
            return False
    
    def generate_document(self, result, template_name=None):
        """
        基于模板生成文档
        :param result: 六壬排盘结果
        :param template_name: 模板名称
        :return: Document 对象
        """
        # 加载模板
        if template_name:
            self.load_template(template_name)
        
        # 创建文档
        doc = Document()
        
        # 应用格式
        self.apply_format(doc)
        
        # 添加内容
        # 设置标题
        doc.add_heading('大六壬排盘结果', level=1)
        
        # 基本信息
        info_paragraph = doc.add_paragraph()
        info_paragraph.add_run('基本信息').bold = True
        info_paragraph.add_run('\n')
        info_paragraph.add_run(f'月将：{result["月将"]}\n')
        info_paragraph.add_run(f'占时：{result["占时"]}\n')
        info_paragraph.add_run(f'日柱：{result["日柱"]}\n')
        
        # 天地盘图片
        from .docx_exporter import LiuRenExporter
        exporter = LiuRenExporter()
        tiandi_pan_img = exporter.generate_tiandi_pan_image(result['天地盘'], result.get('天将'))
        
        import io
        img_bytes = io.BytesIO()
        tiandi_pan_img.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        
        tiandi_paragraph = doc.add_paragraph()
        tiandi_paragraph.add_run('天地盘').bold = True
        tiandi_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(img_bytes, width=Inches(self.current_format['image']['width']))
        
        # 四课图片
        si_ke_img = exporter.generate_si_ke_image(result['四课'], result['日柱'])
        img_bytes = io.BytesIO()
        si_ke_img.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        
        si_ke_paragraph = doc.add_paragraph()
        si_ke_paragraph.add_run('四课').bold = True
        si_ke_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(img_bytes, width=Inches(self.current_format['image']['width'] * 0.67))  # 4/6 = 0.67
        
        # 三传图片
        san_chuan_img = exporter.generate_san_chuan_image(result['三传'])
        img_bytes = io.BytesIO()
        san_chuan_img.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        
        san_chuan_paragraph = doc.add_paragraph()
        san_chuan_paragraph.add_run('三传').bold = True
        san_chuan_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(img_bytes, width=Inches(self.current_format['image']['width'] * 0.67))
        
        # 禄马贵信息
        lu_ma_gui_paragraph = doc.add_paragraph()
        lu_ma_gui_paragraph.add_run('禄马贵').bold = True
        lu_ma_gui_paragraph.add_run('\n')
        lu_ma_gui_paragraph.add_run(f'禄：{result["禄马贵"]["禄"]}\n')
        lu_ma_gui_paragraph.add_run(f'驿马：{result["禄马贵"]["驿马"]}\n')
        lu_ma_gui_paragraph.add_run(f'贵人：{"、".join(result["禄马贵"]["贵人"])}\n')
        
        return doc
